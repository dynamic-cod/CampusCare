import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_proposal_document():
    doc = docx.Document()

    # Set 1.0 inch standard margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # Base Font Setup (Times New Roman, 12pt)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)

    # Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_run = title_p.add_run("SOFTWARE ENGINEERING PROJECT PROPOSAL\n")
    t_run.font.size = Pt(16)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_run = title_p.add_run("ACADEMIC PROJECT SYNOPSIS")
    sub_run.font.size = Pt(13)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(71, 85, 105)
    title_p.paragraph_format.space_after = Pt(14)

    def add_section(heading):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(heading)
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)
        pBrd = parse_xml(f'<w:pBrd {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="CBD5E1"/></w:pBrd>')
        p._p.get_or_add_pPr().append(pBrd)

    # 1. Project Title
    add_section("1. Project Title")
    doc.add_paragraph("CampusCare: Smart Campus Complaint & Maintenance Management System")

    # 2. Problem Identification / Problem Statement
    add_section("2. Problem Identification / Problem Statement")
    doc.add_paragraph(
        "In modern higher education institutions, facility incident tracking and physical maintenance rely heavily "
        "on manual paper registers, verbal complaints to hostel wardens or facility managers, and unorganized email threads. "
        "This manual situation introduces several critical operational limitations:\n"
        "• Time-Consuming & Error-Prone: Physical registers are slow to query, frequently misplaced, and fail to capture standardized spatial data (such as specific room, wing, or floor numbers).\n"
        "• Duplicate Reporting: When a shared facility (such as a lecture hall air conditioner or laboratory projector) breaks down, multiple students log separate complaints for the identical issue, causing redundant paperwork and technician confusion.\n"
        "• Lack of Tracking & Accountability: Students receive no status updates after submitting a request, leading to unresolved complaints with no identifiable owner.\n"
        "• Unilateral Closures: Complaints are frequently marked resolved by maintenance staff without physical verification or complainant feedback.\n\n"
        "Affected Stakeholders: Students, faculty members, laboratory staff, facility technicians (electricians, plumbers, carpenters), and administrative estate officers."
    )

    # 3. Motivation / Need
    add_section("3. Motivation / Need")
    doc.add_paragraph(
        "• Who Faces the Problem: Students and faculty hindered by malfunctioning campus infrastructure, maintenance technicians overwhelmed with unorganized daily requests, and administrators lacking operational oversight.\n"
        "• Why the Existing Method is Inadequate: Manual registers cannot calculate dynamic resolution deadlines, lack automated escalation for overdue tasks, cannot prevent duplicate filings, and provide zero analytical intelligence regarding recurring infrastructure failure patterns.\n"
        "• What Improvement is Expected: A frictionless, automated web-based management platform enabling instant issue reporting via room QR codes, semantic duplicate detection, strict Service Level Agreement (SLA) enforcement, and closed-loop photo-verified ticket resolutions.\n"
        "• Who Will Benefit: Campus students and faculty gain operational infrastructure; technicians receive structured, prioritized daily queues; and estate administrators gain real-time metrics on maintenance efficiency."
    )

    # 4. Objectives
    add_section("4. Objectives")
    objectives = [
        "To develop a centralized spatial registry modeling campus infrastructure across a four-tier hierarchy (Campus → Building → Floor → Room) integrated with printable room QR codes.",
        "To provide a frictionless interface for students to log maintenance complaints without mandatory account registration using unique tracking tokens.",
        "To implement an NLP-based duplicate detection mechanism (TF-IDF and Cosine Similarity) to curb redundant ticket creation in shared locations.",
        "To establish an automated Service Level Agreement (SLA) tracking and background escalation engine that automatically flags overdue maintenance tasks.",
        "To mandate a closed-loop verification workflow requiring staff to submit visual proof-of-work and students to verify and rate completed resolutions.",
        "To generate operational analytics and downloadable audit reports (CSV/PDF) detailing campus maintenance hotspots and department resolution durations."
    ]
    for obj in objectives:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(obj)

    # 5. Proposed Solution
    add_section("5. Proposed Solution")
    doc.add_paragraph(
        "CampusCare delivers a centralized web-based maintenance helpdesk application that automates the complete incident lifecycle. "
        "Every campus room is assigned a unique, printable QR code. When an issue occurs, a student scans the QR code, opening a pre-filled submission form tied to that exact space.\n"
        "As the user types the issue description, an integrated NLP model (TF-IDF Vectorization and Cosine Similarity) evaluates open tickets in that location. If a matching issue is detected (>= 55% similarity), the system prompts the user to upvote the existing ticket rather than creating a duplicate.\n"
        "Upon submission, category-specific SLA deadlines are calculated. Maintenance technicians manage tickets within role-specific dashboards, log status changes, and must upload photographic proof of resolution before marking tasks resolved. Finally, the reporting student confirms resolution quality, provides a 1-to-5 star rating, or reopens the ticket if repairs are inadequate."
    )

    # 6. Preliminary Functional Requirements
    add_section("6. Preliminary Functional Requirements")
    frs = [
        ("FR1: ", "Allow role-based authentication and authorization for Students, Maintenance Staff, and Administrators."),
        ("FR2: ", "Allow administrators to configure spatial records (Campus, Building, Floor, Room) and export printable QR sticker sheets."),
        ("FR3: ", "Allow students to log complaints frictionlessly using room QR codes and generate secure tracking tokens."),
        ("FR4: ", "Analyze complaint text using NLP to detect semantic duplicates at the same location and enable community upvoting."),
        ("FR5: ", "Automatically calculate SLA deadlines and escalate overdue tickets via automated background routines."),
        ("FR6: ", "Allow maintenance staff to update ticket statuses, log internal notes, and submit mandatory resolution proof images."),
        ("FR7: ", "Enable reporting students to verify completed resolutions, submit satisfaction ratings, or reopen unresolved tickets."),
        ("FR8: ", "Generate administrative dashboard charts and export maintenance activity logs as CSV files.")
    ]
    for code, desc in frs:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        r = bp.add_run(code)
        r.font.bold = True
        bp.add_run(desc)

    # 7. Preliminary Non-Functional Requirements
    add_section("7. Preliminary Non-Functional Requirements")
    nfrs = [
        ("NFR1 (Security): ", "Role-Based Access Control (RBAC) shall prevent unauthorized modifications, and CSRF protection must be enforced across all form submissions."),
        ("NFR2 (Usability): ", "The interface shall be fully responsive across mobile, tablet, and desktop viewports to facilitate on-site mobile QR code scans."),
        ("NFR3 (Performance): ", "Duplicate detection evaluations and database queries shall execute within under 1.5 seconds under standard campus traffic."),
        ("NFR4 (Reliability & Auditability): ", "All status changes must be immutably recorded in a status history table with timestamps and actor IDs to ensure an indisputable audit trail."),
        ("NFR5 (Maintainability): ", "The codebase shall follow a modular Django MVT architecture with decoupled configuration and application layers adhering to PEP 8 standards.")
    ]
    for code, desc in nfrs:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        r = bp.add_run(code)
        r.font.bold = True
        bp.add_run(desc)

    # 8. Scope
    add_section("8. Scope")
    doc.add_paragraph(
        "Included:\n"
        "• Spatial campus asset management (Campus, Building, Floor, Room).\n"
        "• Dynamic room QR code generation and batch PDF export.\n"
        "• Complaint submission, reference tracking, and token-based status checks.\n"
        "• NLP-based duplicate detection and community upvoting.\n"
        "• Dynamic SLA tracking, background priority escalation, and audit logging.\n"
        "• Proof-of-work image upload and student satisfaction feedback loop.\n"
        "• Admin dashboard analytics and CSV data export.\n\n"
        "Not Included:\n"
        "• Procurement, purchase order workflows, and inventory purchasing for replacement spare parts.\n"
        "• Financial accounting, budgeting, and technician payroll management.\n"
        "• External third-party contractor and vendor contract management.\n"
        "• Native compiled mobile applications (the platform is deployed as a fully responsive progressive web app)."
    )

    # 9. Feasibility Study
    add_section("9. Feasibility Study")
    doc.add_paragraph(
        "• Technical Feasibility: The system is built using Python and the Django framework, which provides robust built-in security, authentication, and database ORM features. The Scikit-learn library provides lightweight, memory-efficient TF-IDF and Cosine Similarity computations capable of running on standard server hardware without requiring dedicated GPU infrastructure.\n"
        "• Operational Feasibility: The platform minimizes reporting friction. Students scan room QR codes to access clear forms without mandatory account setups. Maintenance technicians manage tickets from mobile-friendly dashboards, eliminating paper-based logs.\n"
        "• Economic / Resource Feasibility: CampusCare utilizes open-source programming frameworks, public packages, and free development tools (Python, Django, VS Code, Git), resulting in zero software licensing costs. Hosting can be achieved on standard institutional servers or free/low-cost cloud tiers."
    )

    # 10. Proposed Technology Stack
    add_section("10. Proposed Technology Stack")
    tech_table = doc.add_table(rows=6, cols=2)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_rows = [
        ("Front-end", "HTML5, CSS3, JavaScript (ES6+), Tailwind CSS"),
        ("Back-end", "Python 3.10+, Django Web Framework (MVT Architecture)"),
        ("Database", "SQLite3 (Development) / PostgreSQL (Production)"),
        ("NLP & Utilities", "Scikit-Learn (TF-IDF, Cosine Similarity), ReportLab, Pillow, QRCode"),
        ("Development Tools", "Visual Studio Code"),
        ("Version Control", "Git / GitHub")
    ]
    for idx, (comp, tech) in enumerate(tech_rows):
        row = tech_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.2), Inches(4.3)
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0)
        set_cell_margins(c1)
        c0.paragraphs[0].add_run(comp).font.bold = True
        c1.paragraphs[0].add_run(tech)

    # 11. Preliminary Development Plan
    add_section("11. Preliminary Development Plan")
    plan_table = doc.add_table(rows=9, cols=2)
    plan_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    phases = [
        ("Phase 1", "Problem Identification & Stakeholder Requirement Elicitation"),
        ("Phase 2", "Requirements Analysis & Formulation of SRS Document"),
        ("Phase 3", "System Architecture & UI/UX Wireframe Design"),
        ("Phase 4", "Relational Database Schema Design & Normalization"),
        ("Phase 5", "Core Backend Implementation (QR utility, NLP Duplicate Engine, SLA Manager)"),
        ("Phase 6", "System Integration, Verification Workflow, & Role-Based Views"),
        ("Phase 7", "Unit Testing, System Validation, & User Acceptance Testing"),
        ("Phase 8", "Live Deployment & Comprehensive Documentation Submission")
    ]
    for idx, (p_num, p_act) in enumerate(phases):
        row = plan_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(1.5), Inches(5.0)
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0)
        set_cell_margins(c1)
        c0.paragraphs[0].add_run(p_num).font.bold = True
        c1.paragraphs[0].add_run(p_act)

    # 12. Expected Outcome
    add_section("12. Expected Outcome")
    doc.add_paragraph(
        "The completed project will provide a production-ready, automated web-based campus facility management system that eliminates manual physical registers. It provides:\n"
        "1. A centralized web platform with secure role-based portals for students, technicians, and administrators.\n"
        "2. Batch printable PDF sheets of room-specific QR codes for campus-wide physical deployment.\n"
        "3. An intelligent duplicate detection filter curbing redundant complaints in shared spaces.\n"
        "4. A transparent, photo-verified resolution and rating workflow ensuring technician accountability.\n"
        "5. Interactive administrative dashboards displaying maintenance efficiency and infrastructure hotspots."
    )

    # 13. References
    add_section("13. References")
    refs = [
        "[1] R. S. Pressman and B. R. Maxim, Software Engineering: A Practitioner's Approach, 9th ed. New York, NY, USA: McGraw-Hill Education, 2020.",
        "[2] A. Silberschatz, H. F. Korth, and S. Sudarshan, Database System Concepts, 7th ed. New York, NY, USA: McGraw-Hill Education, 2019.",
        "[3] Django Software Foundation, \"Django Documentation,\" Django Project, 2024. [Online]. Available: https://docs.djangoproject.com/.",
        "[4] F. Pedregosa et al., \"Scikit-learn: Machine Learning in Python,\" Journal of Machine Learning Research, vol. 12, pp. 2825–2830, 2011.",
        "[5] IEEE Computer Society, \"IEEE Standard for Information Technology—Systems Design—Software Design Descriptions,\" IEEE Std 1016-2009, 2009."
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(ref)

    # Signatures Table
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s0, s1 = sig_table.rows[0].cells[0], sig_table.rows[0].cells[1]
    s0.width, s1.width = Inches(3.25), Inches(3.25)
    s0.paragraphs[0].add_run("_____________________________\nStudent Candidate Signature\nDate: ____ / ____ / 2026").font.size = Pt(10)
    s1_p = s1.paragraphs[0]
    s1_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    s1_p.add_run("_____________________________\nProject Guide / Faculty Approval\nDate: ____ / ____ / 2026").font.size = Pt(10)

    filename = "CampusCare_Formal_Synopsis.docx"
    doc.save(filename)
    print(f"Generated: {filename}")

if __name__ == "__main__":
    create_proposal_document()
